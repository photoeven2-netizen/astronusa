import openpyxl
import json

excel_path = r'C:\Users\Aleecia\Documents\Data_Kabupaten_Kota_Di Dunia.xlsx'
wb = openpyxl.load_workbook(excel_path)

# Standard Country Centroid Coordinates & Timezones
COUNTRY_DEFAULTS = {
  'Indonesia': { 'lat': -2.5489, 'lon': 118.0149, 'tz': 'Asia/Jakarta', 'code': 'ID' },
  'Malaysia': { 'lat': 4.2105, 'lon': 101.9758, 'tz': 'Asia/Kuala_Lumpur', 'code': 'MY' },
  'Singapura': { 'lat': 1.3521, 'lon': 103.8198, 'tz': 'Asia/Singapore', 'code': 'SG' },
  'Thailand': { 'lat': 15.8700, 'lon': 100.9925, 'tz': 'Asia/Bangkok', 'code': 'TH' },
  'Vietnam': { 'lat': 14.0583, 'lon': 108.2772, 'tz': 'Asia/Ho_Chi_Minh', 'code': 'VN' },
  'Filipina': { 'lat': 12.8797, 'lon': 121.7740, 'tz': 'Asia/Manila', 'code': 'PH' },
  'Brunai': { 'lat': 4.5353, 'lon': 114.7277, 'tz': 'Asia/Brunei', 'code': 'BN' },
  'Laos': { 'lat': 19.8563, 'lon': 102.4955, 'tz': 'Asia/Vientiane', 'code': 'LA' },
  'Myanmar': { 'lat': 21.9162, 'lon': 95.9560, 'tz': 'Asia/Yangon', 'code': 'MM' },
  'Timor Leste': { 'lat': -8.8742, 'lon': 125.7275, 'tz': 'Asia/Dili', 'code': 'TL' },
  'ChinaTiongkok': { 'lat': 35.8617, 'lon': 104.1954, 'tz': 'Asia/Shanghai', 'code': 'CN' },
  'Japan': { 'lat': 36.2048, 'lon': 138.2529, 'tz': 'Asia/Tokyo', 'code': 'JP' },
  'Korea Selatan': { 'lat': 35.9078, 'lon': 127.7669, 'tz': 'Asia/Seoul', 'code': 'KR' },
  'Korea Utara': { 'lat': 40.3399, 'lon': 127.5101, 'tz': 'Asia/Pyongyang', 'code': 'KP' },
  'Mongolia': { 'lat': 46.8625, 'lon': 103.8467, 'tz': 'Asia/Ulaanbaatar', 'code': 'MN' },
  'Taiwan': { 'lat': 23.6978, 'lon': 120.9605, 'tz': 'Asia/Taipei', 'code': 'TW' },
  'Indiia': { 'lat': 20.5937, 'lon': 78.9629, 'tz': 'Asia/Kolkata', 'code': 'IN' },
  'Pakistan': { 'lat': 30.3753, 'lon': 69.3451, 'tz': 'Asia/Karachi', 'code': 'PK' },
}

WITA_PROVINCES = {
    'Bali', 'Nusa Tenggara Barat', 'Nusa Tenggara Timur',
    'Kalimantan Selatan', 'Kalimantan Timur', 'Kalimantan Utara',
    'Sulawesi Utara', 'Sulawesi Tengah', 'Sulawesi Selatan', 'Sulawesi Tenggara', 'Gorontalo', 'Sulawesi Barat'
}

WIT_PROVINCES = {
    'Maluku', 'Maluku Utara', 'Papua', 'Papua Barat', 'Papua Selatan', 'Papua Tengah', 'Papua Pegunungan', 'Papua Barat Daya'
}

EXACT_COORDS = {
  'kota tangerang': { 'lat': -6.1783, 'lon': 106.6319 },
  'kabupaten tangerang': { 'lat': -6.1731, 'lon': 106.4673 },
  'kota tangerang selatan': { 'lat': -6.2886, 'lon': 106.7179 },
  'kabupaten tabanan': { 'lat': -8.5411, 'lon': 115.1226 },
  'kota denpasar': { 'lat': -8.6705, 'lon': 115.2126 },
  'kabupaten klaten': { 'lat': -7.7058, 'lon': 110.6064 },
  'kabupaten tuban': { 'lat': -6.8973, 'lon': 112.0649 },
  'kota surabaya': { 'lat': -7.2575, 'lon': 112.7521 },
  'kota bandung': { 'lat': -6.9175, 'lon': 107.6191 },
  'kuala lumpur': { 'lat': 3.1390, 'lon': 101.6869 },
  'tokyo': { 'lat': 35.6762, 'lon': 139.6503 },
  'seoul': { 'lat': 37.5665, 'lon': 126.9780 },
  'beijing': { 'lat': 39.9042, 'lon': 116.4074 },
  'shanghai': { 'lat': 31.2304, 'lon': 121.4737 },
  'bangkok': { 'lat': 13.7563, 'lon': 100.5018 },
}

cities = []
id_counter = 1

for sheetname in wb.sheetnames:
    if sheetname == 'Sheet18': continue
    sheet = wb[sheetname]
    cDefaults = COUNTRY_DEFAULTS.get(sheetname, { 'lat': 0.0, 'lon': 0.0, 'tz': 'Asia/Jakarta', 'code': 'XX' })

    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0: continue
        if not row or not any(row): continue

        # Parse fields based on sheet layout
        prov = ''
        fullName = ''

        if sheetname == 'Indonesia':
            prov = str(row[0] or '').strip()
            jenis = str(row[1] or '').strip()
            nama = str(row[2] or '').strip()
            fullName = f'{jenis} {nama}' if jenis in ['Kabupaten', 'Kota'] else nama
        elif sheetname == 'Malaysia':
            prov = str(row[0] or '').strip()
            jenis = str(row[1] or '').strip()
            nama = str(row[2] or '').strip()
            fullName = nama
        elif sheetname == 'Singapura':
            prov = str(row[0] or '').strip()
            nama = str(row[1] or '').strip()
            fullName = nama
        elif sheetname == 'Thailand':
            prov = str(row[1] or '').strip()
            nama = str(row[3] or row[1] or '').strip()
            fullName = nama
        elif sheetname == 'Vietnam':
            prov = str(row[0] or '').strip()
            nama = str(row[1] or '').strip()
            fullName = nama
        elif sheetname == 'Filipina':
            prov = str(row[0] or '').strip()
            nama = str(row[1] or '').strip()
            fullName = nama
        elif sheetname in ['Brunai', 'Laos', 'Myanmar', 'Timor Leste', 'ChinaTiongkok', 'Japan', 'Korea Selatan', 'Korea Utara', 'Mongolia', 'Taiwan', 'Indiia', 'Pakistan']:
            prov = str(row[0] or '').strip()
            nama = str(row[1] or row[2] or '').strip()
            fullName = nama

        if not fullName or fullName == 'None': continue

        # Determine Timezone for Indonesia
        tz = cDefaults['tz']
        if sheetname == 'Indonesia':
            for w in WITA_PROVINCES:
                if w.lower() in prov.lower():
                    tz = 'Asia/Makassar'
                    break
            for w in WIT_PROVINCES:
                if w.lower() in prov.lower():
                    tz = 'Asia/Jayapura'
                    break

        lowerFull = fullName.lower()
        exact = EXACT_COORDS.get(lowerFull)

        lat = exact['lat'] if exact else cDefaults['lat']
        lon = exact['lon'] if exact else cDefaults['lon']

        cities.append({
            'id': id_counter,
            'name': fullName,
            'province': prov,
            'regionCode': f"{cDefaults['code']}-{(id_counter % 90) + 10}",
            'latitude': lat,
            'longitude': lon,
            'timeZone': tz,
            'population': 300000
        })
        id_counter += 1

outPath = r'E:\birth chart app\apps\astroguide\assets\data\indonesia_cities.json'
with open(outPath, 'w', encoding='utf-8') as f:
    json.dump(cities, f, ensure_ascii=False, indent=2)

print(f'Successfully built {len(cities)} worldwide city records across 18 countries!')
